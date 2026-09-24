"""Read the reconciled workbook without repairing or certifying its contents."""
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
import re
from zipfile import BadZipFile

import openpyxl
import pandas as pd

LAYOUT = json.loads(Path(__file__).with_name("workbook_layout.json").read_text())
PART_CODES = {"part_number", "part_number_normalized", "oem_equivalent_number"}
PART_REFERENCES = PART_CODES | {"oem_battery_part_number", "aftermarket_equivalent_part_numbers",
                               "hv_replacement_assembly_pn", "oem_alternator_pn_ref"}
SCIENTIFIC = re.compile(r"[+-]?\d+(?:\.\d+)?[eE][+-]?\d+")


def load_workbook(path, mode="production"):
    """Return preserved rows, separated research/dev candidates, and located findings.

    Candidates are NOT approved for database insertion: later validators are required.
    ERROR quarantines the affected row; structural errors quarantine the whole file.
    """
    if mode not in {"dev", "production"}:
        raise ValueError("mode must be dev or production")
    path = Path(path)
    result = {"component": "loader-only", "mode": mode,
              "sha256": sha256(path.read_bytes()).hexdigest(), "findings": [],
              "research_candidates": {}, "synthetic_candidates": {}, "quarantined": []}
    records, fatal = [], False

    def finding(sheet, row, column, rule, value, message, severity="ERROR"):
        result["findings"].append(dict(sheet=sheet, row=row, column=column,
                                       rule=rule, value=str(value), message=message,
                                       severity=severity))

    book = openpyxl.load_workbook(path, data_only=False)
    try:
        for name, spec in LAYOUT["sheets"].items():
            head, headers = spec["header_row"], spec["headers"]
            if name not in book.sheetnames:
                finding(name, head, "*", "MISSING_SHEET", "", "Required sheet is absent.")
                fatal = True
                continue
            ws = book[name]
            if any(m.min_row <= head <= m.max_row for m in ws.merged_cells.ranges):
                finding(name, head, "*", "MERGED_HEADER", "", "Merged headers reject the file.")
                fatal = True
            actual = [ws.cell(head, c).value for c in range(1, len(headers) + 1)]
            extra = any(ws.cell(head, c).value is not None
                        for c in range(len(headers) + 1, ws.max_column + 1))
            if actual != headers or extra:
                finding(name, head, "*", "HEADER_LAYOUT", actual,
                        "Headers differ from the explicit workbook layout; review the manifest.")
                fatal = True
            if "data_blocks" in spec:
                rows = {r for a, b in spec["data_blocks"] for r in range(a, b + 1)}
                for r in range(head + 1, ws.max_row + 1):
                    if r not in rows and re.fullmatch(r"S\d+", str(ws.cell(r, 1).value or "")):
                        finding(name, r, "id", "SOURCE_BLOCK_DRIFT", ws.cell(r, 1).value,
                                "Source row lies outside the explicitly registered data blocks.")
                        fatal = True
            else:
                rows = range(head + 1, ws.max_row + 1)
            frame = pd.read_excel(path, sheet_name=name, header=None,
                                  dtype=str, keep_default_na=False, engine="openpyxl")
            for r in sorted(rows):
                cells = [ws.cell(r, c) for c in range(1, len(headers) + 1)]
                if not any(c.value is not None for c in cells):
                    continue
                if any(ws.cell(r, c).value is not None for c in range(len(headers) + 1, ws.max_column + 1)):
                    finding(name, r, "*", "UNHEADED_DATA", "", "Record contains data beyond the declared columns.")
                values, formulas = {}, {}
                for c, column in enumerate(headers, 1):
                    raw = cells[c - 1]
                    value = str(frame.iat[r - 1, c - 1]) if r <= len(frame) and c <= frame.shape[1] else ""
                    if raw.data_type == "f":
                        formulas[column] = raw.value
                        literal = re.fullmatch(r'="([^"\n]*)"', raw.value)
                        if raw.value.upper() in {"=TRUE()", "=FALSE()"}:
                            value = raw.value[1:-2].upper()
                        elif literal:
                            value = literal[1]
                        else:
                            value = raw.value
                            finding(name, r, column, "UNTRUSTED_FORMULA", value,
                                    "Only literal string/boolean formulas are accepted; no formula is evaluated.")
                    elif isinstance(raw.value, str):
                        # pandas may interpret Excel error cells as NaN; preserve the source.
                        value = raw.value
                    values[column] = value
                    machine = column in PART_CODES or column == spec["key"] or column.endswith("_id")
                    if column in {"brand_name", "submodel"} and value != value.strip():
                        finding(name, r, column, "CODE_WHITESPACE", value, "Identity label has boundary whitespace; no trim applied.")
                    if machine and value:
                        if value != value.strip() or "\u00a0" in value or "\u200b" in value:
                            finding(name, r, column, "CODE_WHITESPACE", value, "Review whitespace; value was not trimmed.")
                        if any(ord(ch) > 127 for ch in value):
                            finding(name, r, column, "NON_ASCII_CODE", value, "Review possible Unicode homoglyph; no substitution made.")
                        if isinstance(raw.value, (date, datetime)) or raw.is_date:
                            finding(name, r, column, "DATE_COERCED_CODE", value, "Excel stored a date in a code field.")
                        if SCIENTIFIC.fullmatch(value) or (raw.data_type == "n" and "E+" in raw.number_format.upper()):
                            finding(name, r, column, "SCIENTIFIC_CODE", value, "Scientific notation is ambiguous in identifiers.")
                        if raw.data_type == "n" and len(str(raw.value).replace(".", "").replace("-", "")) > 15:
                            finding(name, r, column, "NUMERIC_PRECISION", value, "Excel numeric identifiers may have lost precision.")
                        if len(value) > 128 or value.endswith(("...", "…")):
                            finding(name, r, column, "CODE_TRUNCATION_RISK", value, "Code exceeds the reviewed limit or ends in an ellipsis.")
                    if column in PART_REFERENCES - PART_CODES and value:
                        if isinstance(raw.value, (date, datetime)) or raw.is_date:
                            finding(name, r, column, "DATE_COERCED_CODE", value, "Excel stored a date in a part reference field.")
                        if SCIENTIFIC.fullmatch(value):
                            finding(name, r, column, "SCIENTIFIC_CODE", value, "Scientific notation is ambiguous in part references.")
                    if column in PART_REFERENCES and raw.value is not None and raw.data_type == "n":
                        finding(name, r, column, "PART_CODE_NOT_TEXT", value,
                                "Numeric part code may have lost leading zeros; restore from its source.")
                    if len(value) >= 32767:
                        finding(name, r, column, "EXCEL_TEXT_LIMIT", value[:80], "Cell reached Excel's text limit; verify the original source.")
                    if re.fullmatch(r"[+-]?\d+,\d+", value):
                        finding(name, r, column, "LOCALE_NUMBER", value, "Comma-separated number is ambiguous; do not normalize silently.")
                    if raw.data_type == "e":
                        finding(name, r, column, "EXCEL_ERROR", value, "Excel error cell cannot be ingested.")
                if (values.get("battery_entry_id") == "B-PENDING"
                        and values.get("verification_status") == "Pending"
                        and not values.get("linked_vcdb_vehicle_ids (Sheet 1)")):
                    finding(name, r, spec["key"], "PENDING_TEMPLATE", "B-PENDING",
                            "Unlinked template excluded from candidate records.", "INFO")
                    continue
                if not values[spec["key"]]:
                    finding(name, r, spec["key"], "MISSING_KEY", "", "Nonempty record has no key.")
                if name == "01b_sources" and not re.fullmatch(r"S\d+", values["id"]):
                    finding(name, r, "id", "SOURCE_KEY", values["id"], "Expected a registered source code S followed by digits.")
                if name == "01_vehicles":
                    allowed = LAYOUT["allowed_submodels"].get(values["make"] + "|" + values["model"], [])
                    if values["submodel"] not in allowed:
                        finding(name, r, "submodel", "TRIM_DRIFT", values["submodel"], "Trim is outside the reviewed exact-value list.")
                source = values.get("data_source_id", "")
                if source and not re.fullmatch(r"\d+", source):
                    finding(name, r, "data_source_id", "SOURCE_ID_INTEGER", source, "Source ID must be an integer token.")
                synthetic = bool(re.fullmatch(r"0*999", source.strip()))
                if synthetic:
                    finding(name, r, "data_source_id", "SYNTHETIC_SOURCE", source,
                            "Source 999 is restricted to a separate dev partition.",
                            "INFO" if mode == "dev" else "ERROR")
                records.append(dict(sheet=name, row=r, values=values, formulas=formulas,
                                    partition="synthetic" if synthetic else "research"))
    finally:
        book.close()

    # Preserve all duplicate records for review; never deduplicate or merge across brands.
    seen, parts, skeletons = {}, {}, defaultdict(list)
    synthetic_ids = {rec["values"]["vcdb_vehicle_id"] for rec in records
                     if rec["sheet"] == "01_vehicles" and rec["partition"] == "synthetic"}
    for rec in records:
        name, r, v = rec["sheet"], rec["row"], rec["values"]
        key = LAYOUT["sheets"][name]["key"]
        identity = (name, v[key])
        if identity in seen:
            for duplicate in (seen[identity], rec):
                finding(name, duplicate["row"], key, "DUPLICATE_KEY", v[key], "Duplicate record key; both records quarantined.")
        seen[identity] = rec
        linkcol = "linked_vcdb_vehicle_ids (Sheet 1)" if name == "02_battery_specifications" else "local_sample_vehicle_ref"
        links = set(re.findall(r"\d+", v.get(linkcol, "")))
        if links & synthetic_ids:
            rec["partition"] = "synthetic"
            finding(name, r, linkcol, "SYNTHETIC_VEHICLE_LINK", v[linkcol],
                    "Record depends on a source-999 vehicle; retained in dev partition only.",
                    "INFO" if mode == "dev" else "ERROR")
        if name == "02_real_parts":
            brand, pn = v["brand_name"], v["part_number"]
            brand_key = v["brand_aaia_id"] or brand.casefold()
            identity = (brand_key, v["part_number_normalized"] or pn)
            if identity in parts:
                for duplicate in (parts[identity], rec):
                    finding(name, duplicate["row"], "part_number", "DUPLICATE_BRAND_PART", pn, "Same brand and part key appears twice.")
            parts[identity] = rec
            skeleton = re.sub(r"[^A-Z0-9]", "", pn.upper()).translate(str.maketrans({"O": "0", "I": "1"}))
            skeletons[(brand_key, skeleton)].append(rec)
    for group in skeletons.values():
        if len({rec["values"]["part_number"] for rec in group}) > 1:
            for rec in group:
                finding(rec["sheet"], rec["row"], "part_number", "CONFUSABLE_PART_CODES",
                        rec["values"]["part_number"], "Same-brand codes differ only by confusable characters or separators; review both.")
    bad_rows = {(f["sheet"], f["row"]) for f in result["findings"] if f["severity"] == "ERROR"}
    for rec in records:
        if fatal or (rec["sheet"], rec["row"]) in bad_rows:
            result["quarantined"].append(rec)
        else:
            result[rec["partition"] + "_candidates"].setdefault(rec["sheet"], []).append(rec)
    result["file_rejected"] = fatal
    result["ok"] = not bad_rows
    result["counts"] = {"inspected_records": len(records), "quarantined": len(result["quarantined"]),
                        "research_candidates": sum(map(len, result["research_candidates"].values())),
                        "synthetic_candidates": sum(map(len, result["synthetic_candidates"].values())),
                        "findings_by_rule": dict(Counter(f["rule"] for f in result["findings"]))}
    return result


def main():
    parser = argparse.ArgumentParser(description="VOLTMAP loader integrity checks (no database writes)")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--mode", choices=["dev", "production"], default="production")
    parser.add_argument("--report", type=Path, help="Write full JSON including preserved candidate/quarantined rows")
    args = parser.parse_args()
    if args.report and args.report.resolve() == args.workbook.resolve():
        parser.error("report path must not overwrite the input workbook")
    try:
        result = load_workbook(args.workbook, args.mode)
    except (OSError, ValueError, KeyError, BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
        parser.exit(2, f"Cannot read workbook: {exc}\n")
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Loader {'PASS' if result['ok'] else 'FAIL'} ({args.mode}); schema/cross-sheet checks still required.")
    print(json.dumps(result["counts"], indent=2))
    for f in result["findings"]:
        print(f"{f['severity']} {f['sheet']}!row {f['row']} [{f['column']}] {f['rule']}: {f['message']}")
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
